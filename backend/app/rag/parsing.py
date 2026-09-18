from __future__ import annotations

import hashlib
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path
from typing import Literal
from zipfile import BadZipFile, ZipFile

import pdfplumber
from docx import Document
from lxml.etree import XMLSyntaxError
from markdown_it import MarkdownIt
from pdfminer.pdfexceptions import PDFException
from pdfplumber.utils.exceptions import PdfminerException


@dataclass(frozen=True)
class Block:
    text: str
    section_path: tuple[str, ...]
    locator: dict[str, int]


@dataclass
class ParsedDocument:
    name: str
    sha256: str
    parser: str
    status: Literal["complete", "partial", "unsupported", "failed"] = "complete"
    blocks: list[Block] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def parse_document(path: Path) -> ParsedDocument:
    if path.stat().st_size > 50 * 1024 * 1024:
        raise ValueError("FILE_TOO_LARGE")
    raw = path.read_bytes()
    suffix = path.suffix.lower()
    result = ParsedDocument(path.name, hashlib.sha256(raw).hexdigest(), f"{suffix}:v1")
    try:
        if suffix in {".md", ".txt"}:
            text = raw.decode("utf-8-sig")
            if "\x00" in text or "\ufffd" in text:
                raise ValueError("INVALID_TEXT")
            if suffix == ".txt":
                result.blocks = [
                    Block(line, (), {"line_start": i, "line_end": i})
                    for i, line in enumerate(text.splitlines(), 1)
                    if line.strip()
                ]
            else:
                _markdown(text, result)
        elif suffix == ".docx":
            _docx(raw, result)
        elif suffix == ".pdf":
            _pdf(raw, result)
        else:
            result.status = "unsupported"
            result.warnings.append("UNSUPPORTED_FORMAT")
    except (
        ValueError,
        UnicodeError,
        BadZipFile,
        PDFException,
        PdfminerException,
        XMLSyntaxError,
        KeyError,
    ):
        result.status = "failed"
        result.blocks.clear()
        result.warnings.append("INVALID_OR_DAMAGED_DOCUMENT")
    if not result.blocks and result.status == "complete":
        result.status = "unsupported"
        result.warnings.append("EMPTY_TEXT")
    elif result.warnings and result.status == "complete":
        result.status = "partial"
    return result


def _markdown(text: str, result: ParsedDocument) -> None:
    tokens = MarkdownIt("commonmark").parse(text)
    headings: list[tuple[int, str]] = []
    lines = text.splitlines()
    for index, token in enumerate(tokens):
        if token.type == "heading_open":
            level = int(token.tag[1:])
            while headings and headings[-1][0] >= level:
                headings.pop()
            headings.append((level, tokens[index + 1].content))
        if token.type == "html_block":
            result.warnings.append(f"HTML_NOT_INTERPRETED:{token.map}")
        if token.type not in {"inline", "fence", "code_block", "html_block"}:
            continue
        if token.map and token.content.strip():
            start, end = token.map
            result.blocks.append(
                Block(
                    "\n".join(lines[start:end]),
                    tuple(h[1] for h in headings),
                    {"line_start": start + 1, "line_end": end},
                )
            )


def _docx(raw: bytes, result: ParsedDocument) -> None:
    with ZipFile(BytesIO(raw)) as archive:
        entries = archive.infolist()
        if len(entries) > 10000 or sum(e.file_size for e in entries) > 100 * 1024 * 1024:
            raise ValueError("UNSAFE_ARCHIVE")
        if "word/document.xml" not in archive.namelist():
            raise ValueError("INVALID_DOCX")
    document = Document(BytesIO(raw))
    if document.tables or document.inline_shapes:
        result.warnings.append("TABLES_OR_IMAGES_NOT_EXTRACTED")
    heading: tuple[str, ...] = ()
    for i, paragraph in enumerate(document.paragraphs, 1):
        if not paragraph.text.strip():
            continue
        if paragraph.style and paragraph.style.name.startswith("Heading"):
            heading = (paragraph.text,)
        result.blocks.append(Block(paragraph.text, heading, {"paragraph": i}))


def _pdf(raw: bytes, result: ParsedDocument) -> None:
    if not raw.startswith(b"%PDF-"):
        raise ValueError("INVALID_PDF")
    with pdfplumber.open(BytesIO(raw)) as pdf:
        if len(pdf.pages) > 500:
            raise ValueError("PAGE_LIMIT")
        for i, page in enumerate(pdf.pages, 1):
            text = page.extract_text() or ""
            if len(text.strip()) < 20:
                result.warnings.append(f"LOW_TEXT_PAGE:{i}")
            if page.find_tables():
                result.warnings.append(f"TABLE_LAYOUT_UNVERIFIED:{i}")
            if text.strip():
                result.blocks.append(Block(text, (), {"page": i}))
