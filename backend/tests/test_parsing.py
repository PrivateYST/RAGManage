from pathlib import Path

from docx import Document

from app.rag.parsing import parse_document


def test_markdown_preserves_evidence_location(tmp_path: Path) -> None:
    path = tmp_path / "rules.md"
    path.write_text("# Booking\n\n## Time\n\nSeven days.\n", encoding="utf-8")
    parsed = parse_document(path)
    block = parsed.blocks[-1]
    assert parsed.status == "complete"
    assert block.section_path == ("Booking", "Time")
    assert block.locator == {"line_start": 5, "line_end": 5}
    assert block.text == "Seven days."


def test_invalid_and_empty_text_are_not_complete(tmp_path: Path) -> None:
    path = tmp_path / "input.txt"
    path.write_bytes(b"\xff\x00")
    assert parse_document(path).status == "failed"
    path.write_bytes(b"")
    assert parse_document(path).status == "unsupported"


def test_raw_html_is_flagged_not_silently_dropped(tmp_path: Path) -> None:
    path = tmp_path / "input.md"
    path.write_text("<script>alert(1)</script>\n", encoding="utf-8")
    parsed = parse_document(path)
    assert parsed.status == "partial"
    assert parsed.warnings


def test_invalid_pdf_is_rejected(tmp_path: Path) -> None:
    path = tmp_path / "input.pdf"
    path.write_bytes(b"not a pdf")
    assert parse_document(path).status == "failed"


def test_damaged_pdf_with_valid_signature(tmp_path: Path) -> None:
    path = tmp_path / "broken.pdf"
    path.write_bytes(b"%PDF-1.7\nnot a document")
    assert parse_document(path).status == "failed"


def test_docx_tables_require_review(tmp_path: Path) -> None:
    path = tmp_path / "rules.docx"
    document = Document()
    document.add_heading("Rules", level=1)
    document.add_paragraph("Booking policy")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Unverified table"
    document.save(str(path))
    parsed = parse_document(path)
    assert parsed.status == "partial"
    assert parsed.blocks[-1].locator == {"paragraph": 2}
    assert parsed.blocks[-1].section_path == ("Rules",)


def test_sample_corpus_preserves_every_source() -> None:
    directory = Path(__file__).resolve().parents[2] / "frontend/docs/knowledgeBase"
    if not directory.is_dir():
        import pytest

        pytest.skip("Local pilot corpus is not distributed in CI")
    paths = sorted(directory.glob("*.md"))
    assert paths
    for path in paths:
        result = parse_document(path)
        assert result.status == "complete", path.name
        lines = path.read_text(encoding="utf-8-sig").splitlines()
        for block in result.blocks:
            start, end = block.locator["line_start"], block.locator["line_end"]
            assert block.text == "\n".join(lines[start - 1 : end])
